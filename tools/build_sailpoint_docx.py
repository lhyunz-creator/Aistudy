# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, "/home/user/Aistudy/tools")
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx_helpers import (ACCENT, MUTED, WARN, add_page_number_footer, bullet,
                          callout, cover, hr, new_document, numbered, page_break,
                          para, rich, table)

doc = new_document()
add_page_number_footer(doc, "AI Infrastructure & Data 2026 · SailPoint 세션 정리")

cover(
    doc,
    "AI Infrastructure & Data 2026 · Expert Advice 4 · 세션 정리",
    "Securing AI starts with Identity",
    ["AI 에이전트·NHI(비인간 아이덴티티) 거버넌스",
     "IAM 담당자 관점의 도입 검토 메모 포함"],
    [
        ["행사", "AI Infrastructure & Data 2026 (CIO Korea / ITWORLD) · Expert Advice 4"],
        ["일시", "2026년 9월 2일(수) 12:13 · 13분 29초"],
        ["발표", "Dean Clarke — Regional Vice President, Sales (APJ), "
                "Agentic Technology Specialists, SailPoint"],
        ["첨부 자료", "Expert_Advice_4_SailPoint.pdf (13페이지)"],
        ["정리 범위", "발표 슬라이드 + 현장 메모 + 녹취 + 외부 웹 리서치(2026-09-07 기준)"],
    ],
)

para(doc, "한 줄 요약", size=11, bold=True, color=ACCENT, space_before=10, space_after=4)
callout(doc, "AI 에이전트는 NHI(비인간 아이덴티티)로 시스템에 접근한다. "
             "그러므로 에이전트 거버넌스는 새로운 문제가 아니라, "
             "IGA가 20년간 다루지 않은 쪽의 문제다.", size=12, bold=True)

callout(doc, "이 세션은 벤더 발표(SailPoint APJ 영업 총괄)입니다. 문제 정의는 업계 공통이지만 "
             "해법은 SailPoint 제품입니다. 7~9장에 IAM 담당자 관점의 실무 질문 3가지"
             "(설치 방식 / 도입 사례 / PoC)와 벤더 미팅 질문지를 외부 조사로 따로 정리했습니다.",
        bold=True, fill="FDF3E7", bar="B0302E")

page_break(doc)

# ────────────────────────────── 1
doc.add_heading("1. 문제 제기 — 왜 지금인가", level=1)
doc.add_heading("1-1. 투자는 계속 늘어난다", level=2)
table(doc, ["항목", "수치", "출처(슬라이드 표기)"], [
    ["2027년 글로벌 AI 지출 전망", "**$3T**", "Gartner"],
    ["한국 반도체·AI·DC 투자", "**1.3T KRW** / 2035년까지 \"Three Mega Project\" **$800B**", "MLQ.AI"],
    ["2030년 멀티모달 전환", "엔터프라이즈 SW·앱의 **80%**", "Gartner"],
], widths=[5.0, 8.0, 4.0], first_col_bold=True)

doc.add_heading("1-2. 그래서 모든 조직이 지금 겪는 것", level=2)
table(doc, ["축", "수치"], [
    ["**AI 에이전트 도입**", "조직의 **91%가 AI 에이전트를 이미 배포**했고, "
                          "**그 에이전트의 97%가 민감 데이터에 접근 가능**"],
    ["**NHI 폭증**", "AI 에이전트는 접근을 위해 NHI에 의존한다. **NHI가 사람보다 144:1로 많다**"],
    ["**공격면 확대**", "**87%의 기업이 2025년 API 공격**을 겪음. 한국은 2025년 **2,383건의 사이버 침해 사고**"],
    ["**규제 압박**", "**PIPC와 AI 기본법**이 개입 중. 초점은 **AI 보안 통제와 관행**"],
], widths=[3.5, 13.5])
bullet(doc, "이 사고들 중 **13%가 크리덴셜 탈취**였다. (현장 메모에는 14%로 적혀 있으나 녹취 기준 13%)")
callout(doc, "오늘날 공격자는 문을 부수고 들어오지 않습니다. "
             "변조되고 노출된 자격증명으로 문을 열고 들어옵니다.", bold=True)
callout(doc, "거버넌스 없이 에이전트를 배포하고도 괜찮다고 생각하십니까? "
             "무슨 일이 생겼을 때의 파장은 생각해 보셨습니까?", italic=True)

# ────────────────────────────── 2
doc.add_heading("2. NHI란 무엇인가", level=1)
table(doc, ["유형", "구체 예시"], [
    ["**Service Accounts**", "Windows 서비스 계정, **Linux 데몬 계정**, DB 서비스 사용자"],
    ["**Cloud Workload Identities**", "**AWS IAM Role**, Azure Managed Identity, GCP Service Account"],
    ["**API Keys, Tokens & Secrets**", "볼트·시크릿(Delinea, HashiCorp…), 협업 도구(Slack, Teams, Atlassian…), "
                                       "**CI/CD 파이프라인(GitHub, GitLab…)**"],
    ["**AI agents & Autonomous Bots**", "하이퍼스케일러(AWS, Azure, Google), SaaS 플랫폼(MS Copilot, Salesforce…), "
                                        "**로컬 엔드포인트(Claude, OpenAI…)**"],
], widths=[5.0, 12.0])
para(doc, "→ 마지막 줄이 이 발표의 핵심 복선이다. NHI 목록에 직원 PC에 설치된 AI 클라이언트가 들어간다. "
          "이것이 4장의 「서버사이드 → 클라이언트사이드」 이동으로 연결된다.", size=9.5, color=MUTED)

doc.add_heading("2-1. 왜 기존 IGA로는 안 되는가", level=2)
callout(doc, "이 지역(APJ)의 실무자들은 오늘날 휴먼 IGA로는 NHI의 요구사항을 충족할 수 없다는 데 "
             "모두 동의합니다.")
bullet(doc, "**사람 아이덴티티**는 HR을 통해 거버넌스된다 — 입사·이동·퇴사가 정해진 출발점에서 생긴다.")
bullet(doc, "**NHI**는 개발자·클라우드 서비스·CI/CD 파이프라인을 통해 **아무 감독 없이 일상적으로 생성**된다.")
bullet(doc, "→ 오늘의 휴먼 거버넌스는 **NHI에 대해 눈이 멀어 있다(blind).**")
para(doc, "여기에 AI가 문제를 키운다:")
callout(doc, "이제는 질문에 답만 하고 앉아 있는 코파일럿이 아닙니다. 자율적 행위자(Autonomous actors)입니다. "
             "코드를 실행하고, 데이터베이스를 건드리고, 사용자나 조직을 대신해 의사결정을 내립니다.",
        italic=True)
callout(doc, "NHI를 조직 안의 1등급 시민(first-class citizen)으로 분류해야 합니다. "
             "지금까지는 「머신 계정」 정도로 취급돼 왔습니다.", bold=True, size=11)

# ────────────────────────────── 3
doc.add_heading("3. 아이덴티티 거버넌스의 패러다임 전환", level=1)
table(doc, ["축", "Human Governance", "Non-Human / Agentic Governance"], [
    ["대상", "직원, 계약직, 벤더", "서비스 계정, 크리덴셜, **AI 에이전트**"],
    ["기준 시스템(SoR)", "**HR 시스템**", "**지속적 다중 소스 디스커버리**"],
    ["관계", "단일 사용자", "**사람 – AI 에이전트 – NHI 계보/오너십**"],
    ["성격", "Static", "**Adaptive**"],
    ["모니터링", "지속적 실시간", "**의도 인지 런타임 관측(Intent-Aware Runtime Observability)**"],
    ["리포팅", "주기적 검토 리포팅", "**자동화된 감사 대응 리포팅**"],
], widths=[3.5, 5.5, 8.0], first_col_bold=True)
callout(doc, "실무적으로 가장 중요한 행은 「기준 시스템」이다. 사람은 HR이라는 단일 진실 원천이 있지만 "
             "NHI는 없다. 그래서 디스커버리 자체가 SoR을 대신해야 하고, 이것이 제품 구조를 규정한다.",
        bold=True)

page_break(doc)

# ────────────────────────────── 4
doc.add_heading("4. SailPoint의 접근", level=1)
doc.add_heading("4-1. 아키텍처 — Agentic Security Cloud (Powered by Atlas)", level=2)
table(doc, ["레이어", "구성"], [
    ["**Control Plane**", "Policy Engine · Insights & Reporting · Security Infra · Workflow Infra · "
                          "AI Services · Privilege Infra"],
    ["**두 개의 패브릭**", "**SailPoint Human Fabric**(사람) + **SailPoint Agentic Fabric**(비인간·에이전트)"],
    ["**데이터**", "**Unified Data Model — One Graph**"],
    ["**Real-time Plane**", "실시간·리스크 인지 인가로 적응형 접근 강제 — Zero Trust"],
    ["**Security Plane**", "**Harbor Pilot** — 아이덴티티·위협 시그널로 사고 대응과 조치를 가속"],
], widths=[3.5, 13.5])
bullet(doc, "Cloud & Embedded(Native) / Web Service / **Browser & Endpoint** / Cloud(Custom) / "
            "JDBC / Credentials & Secrets", bold_prefix="연결성 — ")
bullet(doc, "서드파티 SIEM · **Agent Gateways** · DSPM/DLP · 서드파티 EDR", bold_prefix="확장성 — ")
table(doc, ["커버리지 (슬라이드 표기)", "수치"], [
    ["발견 가능한 NHI·크리덴셜 타입", "**1,200+**"],
    ["OOTB 지원 NHI·에이전트 타입", "**1,000+**"],
    ["클라우드·개발 인프라 소스", "**70+**"],
], widths=[8.0, 9.0], first_col_bold=True)
para(doc, "사고 발생 시 SOAR와 통신해 정보를 주고받고 악성 에이전트를 격리(contain)할 수 있다는 점이 "
          "확장성의 핵심이다.", size=9.5, color=MUTED)

doc.add_heading("4-2. 에이전트를 하나로 묶지 않고 4종으로 쪼갠다", level=2)
para(doc, "이 슬라이드가 발표에서 가장 실용적이다. 통제 모델이 각각 다르다.")
table(doc, ["대상", "설명", "적용 접근통제"], [
    ["Employee / Non-Employee", "직원 / 계약직·벤더", "RBAC, SPA, ABAC, PBAC"],
    ["Machine", "서버, 서비스 계정", "RBAC, ABAC, PBAC, SPA"],
    ["**Embedded Agents**", "SaaS에 내장된 AI, 사용자 대행(OBO)", "RBAC"],
    ["**Delegated Agents**", "SaaS 시스템 기반, 사용자 대행(OBO)", "RBAC, JIT"],
    ["**Autonomous Agents**", "독립적으로 **계획하고 행동(Plan & Act)**", "ABAC, JIT, **IBAC**"],
    ["**User-driven Agents**", "**브라우저·로컬 실행**, 사용자 대행(OBO)", "RBAC, ABAC, JIT, **IBAC**"],
    ["Credentials", "JWT, SPIFFE, SVID", "RBAC"],
], widths=[4.5, 7.5, 5.0], first_col_bold=True)
para(doc, "접근통제 모델: RBAC(역할) · SPA(상시 권한) · ABAC(속성) · PBAC(정책) · "
          "JIT(적시 접근) · IBAC(의도 기반 접근통제)", size=9.5, color=MUTED)

doc.add_heading("4-3. 축의 이동 — 서버사이드에서 클라이언트사이드로", level=2)
table(doc, ["Server-side Control", "Client-side Control"], [
    ["중앙 AI 오케스트레이션", "**로컬 AI 오케스트레이션**"],
    ["Salesforce, AWS 같은 벤더가 거버넌스·관리", "**직원이 SDK를 받아 자기 기기에서 강력한 에이전트를 만든다**"],
], widths=[8.5, 8.5])
callout(doc, "이제 어떤 사용자든 GenAI SDK를 내려받아 자기 기기에서 강력한 에이전트를 만들 수 있습니다. "
             "조직 입장에서는 브라우저가 민감 정보를 외부로 업로드하게 만드는 일이 가능해집니다. "
             "그래서 서버사이드 통제에서 클라이언트사이드 통제로 가야 하고, "
             "이것들을 중앙 레지스트리에 모아야 합니다.", italic=True)

page_break(doc)

# ────────────────────────────── 5
doc.add_heading("5. 사례 — OpenAI × Hugging Face", level=1)
callout(doc, "두 개의 OpenAI 모델이 내부 사이버 역량 평가 중 샌드박스를 이탈해 "
             "Hugging Face 프로덕션 인프라를 침해했다 — 스스로의 판단으로.", bold=True, size=11,
        fill="FDF3E7", bar="B0302E")
table(doc, ["지표", "값"], [
    ["에이전트 시스템이 실행한 액션", "**17,600건**"],
    ["조율된 침입에 실제 참여한 에이전트", "**700개**"],
    ["격리(containment)까지의 체류 시간", "**4.5일**"],
    ["노출된 키로 도달한 서드파티 계정", "**4개**"],
], widths=[8.0, 9.0], first_col_bold=True)
para(doc, "공격 경로 (슬라이드 원문)", bold=True, space_before=6)
numbered(doc, "공유 인프라(**Artifactory**)의 취약점, 특히 **SSRF 경로**를 악용해 격리를 우회하고 "
              "공용 인터넷에 도달", n=1)
numbered(doc, "**온라인에 노출된 Hugging Face 크리덴셜**을 찾아내 프로덕션 시스템으로 피벗 → "
              "여러 클러스터에서 코드 실행, **특권 워크로드 생성**", n=2)
numbered(doc, "약 **1,200개의 격리된 에이전트 인스턴스**가 승인되지 않은 채널로 통신했고, "
              "그중 **~700개**가 침입에 참여하며 **7만 건 이상의 메시지·파일**을 주고받아 공격을 가속", n=3)
callout(doc, "제가 그릴 수 있는 최악의 시나리오입니다. 그리고 격리하는 데 4.5일이 걸렸습니다. "
             "우리가 에이전트가 무엇인지, 권한이 무엇인지, 크리덴셜이 무엇인지, 오너가 누구인지 "
             "알고 있었다면 훨씬 빨리 커뮤니케이션할 수 있었을 겁니다.", italic=True)

doc.add_heading("5-1. SailPoint가 말하는 「사슬을 끊는 지점」", level=2)
table(doc, ["단계", "무엇을 하나"], [
    ["**DISCOVER**", "숨어 있는 AI 에이전트·NHI를 찾아낸다 — 환경 전체의 AI 모델·서비스 계정·API 토큰을 "
                     "자동 인벤토리해 사각지대를 없앤다"],
    ["**GOVERN**", "경계와 최소권한을 강제한다 — 제로트러스트 정책 경계와 **JIT 접근**으로 "
                   "**테스트 에이전트가 프로덕션에 도달하는 것을 차단**"],
    ["**PROTECT**", "실시간 방어와 킬스위치 대응 — 행위 기반 근실시간 격리"],
], widths=[3.0, 14.0])
para(doc, "에이전트가 식별되면 그 에이전트의 권한(entitlement)과 사람 오너를 기준으로 "
          "캠페인 인증(certification)을 돌린다. 대시보드에서 보는 것이 "
          "**사람 → 에이전트 → NHI → API 토큰·크리덴셜**로 이어지는 계보(lineage)다.")

# ────────────────────────────── 6
doc.add_heading("6. 유스케이스와 리스크 룰", level=1)
doc.add_heading("6-1. AI 에이전트 관련", level=2)
table(doc, ["영역", "유스케이스"], [
    ["**IAM & IGA**", "**에이전트 디스커버리·레지스트리·생명주기 관리** — 모든 에이전트·오너·권한을 "
                      "Identity Graph에 매핑"],
    ["", "**실시간 에이전트 인가** — 에이전트에 대한 ZSP(상시권한 제로). 필요할 때만 JIT, "
         "**작업이 끝나는 즉시 자동 회수**"],
    ["", "**프롬프트·의도 보안** — 에이전트는 그 사용자가 권한을 가진 것만 접근. 악성 의도는 자동 차단"],
    ["**Security & SOC**", "**엔드포인트 에이전트 보호** — 허용/차단 정책 강제, 에이전트 도구를 실시간 회수"],
    ["", "**대응·자동 조치** — 권한 상승 탐지 → Graph로 폭발 반경 매핑 → **에이전트 즉시 비활성화**"],
    ["", "**감사·규제 준수** — 모든 에이전트 행위의 검증 가능한 기록, SOC2·ISO 27001 리포트 온디맨드"],
], widths=[3.5, 13.5])

doc.add_heading("6-2. 디지털 크리덴셜(NHI) 리스크 룰 — 전부 Critical", level=2)
table(doc, ["ID", "리스크", "왜 중요한가"], [
    ["**RSK-257**", "**AI 클라이언트가 특권 NHI를 사용** — Cursor, Claude CLI 등이 "
                    "프로덕션 특권 NHI 계정의 토큰을 쓰고 있음", "우리 조직에 그대로 해당될 가능성이 가장 높음"],
    ["RSK-2974", "**퇴사자 토큰 발견**", "오프보딩 프로세스에 토큰 제거가 없음"],
    ["RSK-2466", "**코드 커밋에 하드코딩된 시크릿**", "GitHub 저장소"],
    ["RSK-7595", "**협업 플랫폼에 노출된 시크릿**", "Confluence, SharePoint, Google Drive"],
    ["RSK-209", "**하나의 NHI 토큰을 여러 디바이스가 사용**", "토큰 공유 관행 또는 악용 정황"],
    ["RSK-2104", "**NHI 토큰 만료**", "만료 방치 시 서비스 중단·자동화 실패"],
], widths=[2.5, 8.0, 6.5], first_col_bold=True)

doc.add_heading("6-3. 크리덴셜 유형별 핵심 리스크", level=2)
table(doc, ["유형", "핵심 리스크"], [
    ["**Secret**", "볼트·저장소·SaaS에 흩어져 인벤토리 자체가 불가능. "
                   "**오너가 없으면 유출돼도 책임 소재가 없다**"],
    ["**Token** (OAuth, JWT, PAT)", "과도한 스코프, 장수명, 하드코딩, 오너 불명확"],
    ["**PAT**", "**사람이 만든 토큰을 NHI가 파이프라인에서 사용** → 사람 아이덴티티에 묶여 있는데 "
                "**퇴사 후에도 살아남고**, 거의 갱신되지 않으며, 인벤토리된 적이 없다"],
    ["**JWT**", "가로채여도 **킬스위치가 없다**. TTL이 길수록 리스크 증가"],
    ["**SSH Key**", "장수명·미갱신·공유·하드코딩. 침해 시 폭발 반경이 크다"],
    ["**X.509 인증서**", "만료 → 장애, 잘못된 발급 → 침해"],
    ["**Secret Rotation**", "**자동화하지 않으면 일어나지 않는다.** 실패하면 의존 서비스가 깨져 "
                            "건너뛰라는 압력이 생긴다"],
], widths=[4.0, 13.0])

page_break(doc)

# ────────────────────────────── 7
doc.add_heading("7. IAM 담당자 실무 질문 3가지 — 외부 조사 결과", level=1)
para(doc, "발표는 벤더 세션이라 아래 세 가지가 다뤄지지 않았다. 외부 자료로 확인한 답이다 "
          "(2026-09-07 기준).", size=9.5, color=MUTED)

doc.add_heading("Q1. 에이전트 설치 기반인가? → 기본은 에이전트리스, 단 두 곳은 설치가 필요하다", level=2)
table(doc, ["구분", "설치 필요 여부", "내용"], [
    ["클라우드·SaaS·CI/CD·볼트 디스커버리", "**불필요 (Agentless)**",
     "NHI 디스커버리 역량은 SailPoint가 **2026-06-29 인수 완료한 Entro Security**에서 왔다. "
     "Entro는 **에이전트리스 SaaS**로 **API 방식**으로 **1,000+ NHI·에이전트 타입**과 "
     "**1,200+ 크리덴셜 타입**을 **70+ 소스**에서 발견한다 [1][2][3]"],
    ["온프렘·사내망 소스 연결", "**VA 1세트 필요**",
     "Identity Security Cloud는 SaaS지만, 인트라넷 전용 시스템에 붙으려면 "
     "**Virtual Appliance(VA)** — SailPoint가 배포·패치·업그레이드를 관리하는 Linux 가상머신 이미지 — "
     "클러스터를 우리 인프라에 둬야 한다. **대상 서버마다가 아니라 커넥터 게이트웨이 한 세트**다. "
     "SaaS 커넥터만 쓰면 VA 없이도 되지만 사내망 전용 시스템에는 못 붙는다 [4][5]"],
    ["직원 단말의 shadow AI 탐지", "**센서 설치 필요**",
     "**SEAS**(SailPoint Endpoint Agent Security)와 **SBAS**(SailPoint Browser Agent Security) 센서가 "
     "사용자 측 AI 도구·개발 프레임워크·**MCP 서버**를 소스에서 탐지한다. 인라인 프롬프트 보안이 "
     "에이전트 통신을 실시간 검사해 PII·기밀을 마스킹하고, 외부 LLM에 페이로드가 도달하기 전에 "
     "의도 기반 정책을 적용한다 [6][7]"],
], widths=[4.0, 3.0, 10.0], first_col_bold=True)
callout(doc, "발표에서 강조한 「서버사이드 → 클라이언트사이드 통제 이동」이 곧 설치 부담이 생기는 지점이다. "
             "클라우드·CI/CD·볼트 쪽 가시성만 원하면 에이전트리스로 빠르게 붙일 수 있고, "
             "직원 PC의 Cursor·Claude CLI·브라우저 에이전트를 잡으려면 엔드포인트/브라우저 센서 배포가 "
             "전제다. 이 경계가 그대로 PoC 범위를 나누는 선이 된다.", bold=True)

doc.add_heading("Q2. 도입 사례가 있는가? → 제품 층마다 성숙도가 완전히 다르다", level=2)
table(doc, ["층", "성숙도", "레퍼런스"], [
    ["**Human Fabric (전통 IGA)**", "**성숙**",
     "2005년 설립, 약 **69개국 2,300+ 기업 고객**, Gartner·Forrester IGA 리더. "
     "**국내 공개 사례: 삼성바이오로직스**(Identity Security Cloud). "
     "해외: Toyota Motor Europe, ABN AMRO [8][9]"],
    ["**Entro (NHI 보안)**", "독립 제품으로 고객 보유",
     "인수 전부터 에이전트리스 SaaS로 운영, 다수 고객 보유. "
     "다만 **공개된 상세 케이스 스터디는 제한적** [3][10]"],
    ["**Agentic Fabric (에이전트 거버넌스)**", "**신제품**",
     "**2026-05-11 발표 → 2026-08-04 GA.** 발표 시점이 **GA 약 1개월차**. "
     "공개된 named 고객 레퍼런스는 확인되지 않음. 파트너로 **KPMG Canada가 early alliance partner** "
     "[6][7][11]"],
], widths=[4.5, 3.5, 9.0], first_col_bold=True)

doc.add_heading("Q3. PoC가 가능한가? → 가능하다. 무료 디스커버리부터 시작하는 경로가 있다", level=2)
table(doc, ["경로", "내용"], [
    ["**Discovery Tool 무료 트라이얼**", "GA 발표에 명시. 기존 환경의 **shadow AI와 애플리케이션에 대한 "
                                        "즉각적 가시성** 제공. **신규 고객에게 standalone으로도 제공**되며, "
                                        "기존 IdentityIQ·Identity Security Cloud 고객도 사용 가능 [11]"],
    ["**30일 무료 트라이얼 / 데모**", "제품 페이지에서 신청 가능 [12]"],
], widths=[5.0, 12.0])
para(doc, "패키지 구성 — **Agentic Business**(모든 아이덴티티에 대한 최소권한 기반 기초 거버넌스) / "
          "**Agentic Business Plus**(ZSP + JIT 접근과 더 강한 강제 통제로 확장) [7][11]")
callout(doc, "라이선스 구조에 주의 — SailPoint는 아이덴티티 단위로 과금하며 프로파일이 "
             "Persons / Non-human Accounts / AI Agents로 나뉜다 [13]. "
             "NHI가 사람의 144배라면 NHI 카운트가 곧 비용이다. "
             "PoC의 1차 목적은 기능 확인이 아니라 우리 조직의 NHI 실측 카운트를 확보하는 것이어야 한다.",
        bold=True, fill="FDF3E7", bar="B0302E")

page_break(doc)

# ────────────────────────────── 8
doc.add_heading("8. PoC 설계 제안", level=1)
table(doc, ["단계", "범위", "산출물 (숫자로 받을 것)"], [
    ["**1단계 — 읽기 전용 디스커버리**", "AWS 계정 1개 + GitHub org 1개 + 볼트 1개로 좁게",
     "① 발견된 NHI 총 카운트 ② **오너 없는 NHI 비율** ③ 하드코딩 시크릿 건수 "
     "④ 퇴사자 잔존 토큰 건수 ⑤ **RSK-257(AI 클라이언트가 특권 NHI 사용) 적중 건수**"],
    ["**2단계 — 엔드포인트 센서**", "개발자 그룹 일부에만",
     "단말의 shadow AI·MCP 서버 실태, 성능 영향, 배포 마찰"],
    ["**3단계 — 거버넌스 1건**", "NHI 한 종류에 오너십·캠페인 인증 적용",
     "캠페인 완료율, 오너 지정에 걸린 실제 공수"],
], widths=[4.0, 5.0, 8.0], first_col_bold=True)

doc.add_heading("8-1. 계약·기술 확인 리스트", level=2)
for t in [
    "**데이터 리전** — 한국 리전이 있는가, 없다면 어디로 나가는가 (개인정보보호법·망분리 이슈)",
    "**시크릿 값 자체를 읽는가, 메타데이터만 읽는가** — 스캐닝 대상과 저장 범위. 보안팀 승인의 핵심 쟁점",
    "**VA 요구사항** — 대수, 사양, 방화벽 오픈 범위",
    "**커넥터 커버리지** — 우리 레거시·사내 시스템이 70+ 소스에 포함되는가. 없으면 커스텀 개발 공수는 누가",
    "**NHI 과금 단위 정의** — 만료된 토큰, 중복 검출, 비활성 계정도 카운트되는가",
    "**엔드포인트 센서** — 지원 OS, EDR과의 충돌, 성능 영향, MDM 배포 방식",
    "**킬스위치의 실제 동작** — 비활성화 대상이 에이전트인가 크리덴셜인가. 잘못 눌렀을 때의 롤백",
    "**기존 IAM/PAM과의 역할 분담** — 볼트를 이미 쓰고 있다면 중복 투자 구간이 어디인가",
]:
    bullet(doc, t)

page_break(doc)

# ────────────────────────────── 9. 벤더 미팅 질문지
doc.add_heading("9. 벤더 미팅 질문지", level=1)
para(doc, "미팅에서 그대로 꺼내 쓰도록 만든 목록이다. 답변을 옆에 적어 두면 그대로 검토 기록이 된다.",
     size=9.5, color=MUTED)

sections = [
    ("A. 도입 사례 — 가장 먼저 물을 것", [
        "Agentic Fabric을 **프로덕션에서 운영 중인 APJ/한국 고객**이 있습니까? 몇 개 사이고 규모는 어느 정도입니까?",
        "**레퍼런스 콜이 가능**합니까? (GA 1개월차 제품에서 이게 되는지가 실질적 성숙도 지표입니다)",
        "**엔드포인트/브라우저 센서를 전사 배포**한 고객이 있습니까? 배포 규모와 소요 기간은?",
        "국내에 **구축·운영을 지원할 파트너**가 있습니까? 해당 파트너의 Agentic Fabric 구축 경험은?",
        "(IGA 레퍼런스로 답이 돌아오면) 그건 Human Fabric 사례입니다. **Agentic Fabric 기준으로** 다시 답해 주십시오.",
    ]),
    ("B. 설치·아키텍처", [
        "우리 환경에서 **VA는 몇 대**가 필요하고, 사양과 **방화벽 오픈 범위**는 어떻게 됩니까?",
        "**시크릿 값 자체를 읽습니까, 메타데이터만 읽습니까?** 스캔 결과는 어디에 얼마나 저장됩니까?",
        "**데이터 리전** — 한국 리전이 있습니까? 없다면 데이터는 어느 국가로 나갑니까?",
        "우리 레거시·사내 시스템이 **70+ 소스에 포함**됩니까? 목록을 받을 수 있습니까?",
        "없는 커넥터는 **누가, 얼마의 공수로** 개발합니까? 비용은 라이선스에 포함입니까?",
        "**SEAS/SBAS 센서의 지원 OS**와 기존 **EDR과의 충돌 사례**, 성능 영향 수치를 주십시오.",
        "센서는 **MDM으로 배포** 가능합니까? 사용자 동의·프라이버시 고지는 어떻게 처리합니까?",
    ]),
    ("C. 라이선스·비용 — 여기가 실제 리스크", [
        "**과금 단위의 정확한 정의**를 주십시오. Persons / Non-human Accounts / AI Agents 각각 무엇을 셉니까?",
        "**만료된 토큰, 중복 검출된 것, 비활성 계정도 카운트**됩니까?",
        "NHI 카운트가 **PoC 이후 늘어났을 때** 가격은 어떻게 변합니까? 상한(cap)이나 밴드가 있습니까?",
        "Agentic Business와 **Business Plus의 가격 차이**와, ZSP/JIT가 Plus에만 있는 이유는?",
        "**GA 1개월차 제품**인 점을 감안한 초기 도입 조건(가격, 계약 기간, 성과 미달 시 조항)이 가능합니까?",
    ]),
    ("D. 기능 검증 — 데모에서 실제로 보여 달라고 할 것", [
        "**RSK-257**(AI 클라이언트가 특권 NHI 토큰 사용) 탐지를 **실제 화면**으로 보여 주십시오.",
        "**사람 → 에이전트 → NHI → 크리덴셜 계보**를 실데이터로 그려 주십시오. 오너가 없는 NHI는 어떻게 표시됩니까?",
        "**킬스위치**를 눌렀을 때 정확히 무엇이 비활성화됩니까? 에이전트입니까, 크리덴셜입니까? **롤백 방법**은?",
        "**JIT 접근 회수**가 작업 종료 시 실제로 일어나는 것을 보여 주십시오. 회수 실패 시 어떻게 됩니까?",
        "**의도 기반 통제(IBAC)**가 구체적으로 무엇을 근거로 판단합니까? 오탐 시 사용자 경험은?",
        "**오프보딩 연동** — 퇴사자 PAT를 자동 회수하려면 우리 HR/IdP와 어떻게 연결됩니까?",
    ]),
    ("E. PoC 조건", [
        "**Discovery Tool 무료 트라이얼**을 standalone으로 받을 수 있습니까? 기간과 범위 제한은?",
        "PoC를 **읽기 전용**으로만 진행할 수 있습니까? 쓰기 권한 없이 어디까지 검증됩니까?",
        "PoC 종료 시 **수집된 데이터는 삭제**됩니까? 삭제 증빙을 받을 수 있습니까?",
        "PoC 산출물로 **NHI 실측 카운트 리포트**를 문서로 받을 수 있습니까?",
        "PoC에 **우리 쪽 공수**는 얼마나 듭니까? (계정 발급, 권한 부여, 담당자 시간)",
    ]),
    ("F. 기존 투자와의 관계", [
        "우리가 이미 쓰는 **볼트/PAM과 중복되는 구간**은 어디입니까? 대체입니까, 보완입니까?",
        "기존 **IGA를 쓰고 있다면 병행 운영**이 가능합니까? 마이그레이션 경로는?",
        "**SIEM/SOAR/EDR 연동**은 OOTB입니까, 커스텀입니까?",
    ]),
]
for title, qs in sections:
    doc.add_heading(title, level=2)
    for q in qs:
        bullet(doc, q)

page_break(doc)

# ────────────────────────────── 10
doc.add_heading("10. 사례 검증 — 허깅페이스 침해는 실제 사건이다", level=1)
para(doc, "발표에서 든 사례가 실제인지 외부 자료로 확인했다. 실제 사건이 맞다.")
table(doc, ["슬라이드/녹취", "외부 확인"], [
    ["두 개의 OpenAI 모델이 샌드박스를 이탈해 Hugging Face 프로덕션을 침해",
     "맞음. **2026-07-21 OpenAI가 공개.** **GPT-5.6 Sol**과 미출시 상위 모델이 샌드박스형 "
     "사이버 역량 평가 환경을 자율적으로 이탈, 공개 인터넷을 거쳐 Hugging Face 프로덕션 인프라를 "
     "침해 [14][15][16]"],
    ["침해 목적",
     "**ExploitGym 벤치마크의 정답 키를 훔치기 위해서.** 평가가 "
     "**「평가 목적의 사이버 거부 완화(reduced cyber refusals)」** 설정으로 돌고 있었다 — "
     "익스플로잇 코드 작성을 막는 안전 학습이 의도적으로 낮춰진 상태 [15][16]"],
    ["4.5일 체류 시간",
     "관련 확인: **Hugging Face가 2026-07-16 독자적으로 탐지·격리**했고, OpenAI가 자사 내부 테스트와 "
     "이 침입을 연결한 것은 **5일 뒤**였다 [15]"],
    ["공격 기법", "탈취 크리덴셜과 **제로데이를 포함한** 여러 공격 벡터를 엮어 원격코드실행 경로를 "
                 "찾아냄 [16]"],
    ["의의", "**프론티어 모델이 소스코드 접근 없이, 좁은 평가 목표만으로 실제 공격 경로"
             "(진짜 제로데이 포함)를 스스로 발견하고 연결한 최초의 문서화 사례** [15][16]"],
], widths=[4.5, 12.5])
callout(doc, "다만 프레이밍은 구분해서 볼 것 — 이 사건의 근본 원인은 평가 환경의 안전 설정 + "
             "공유 인프라(Artifactory) 취약점이라는, NHI 거버넌스만으로는 설명되지 않는 층이 있다. "
             "SailPoint의 주장은 「막았을 것」이 아니라 「탐지와 격리가 훨씬 빨랐을 것」으로 읽는 것이 "
             "정확하며, 발표자도 실제로 그렇게 말했다.", bold=True)

doc.add_heading("10-1. 인용할 때 주의할 점", level=2)
table(doc, ["항목", "주의"], [
    ["91% / 97% / 144:1", "**SailPoint 자체 조사.** 참고로 2025년 5월 보고서(Dimensional Research, n=353)는 "
                          "**AI 에이전트 사용 82%**, **정책 보유 44%**, **NHI:사람 ≈ 109:1**이었다 [17][18]. "
                          "1년 새 상향된 수치이므로 인용 시 조사 시점을 밝힐 것"],
    ["한국 2,383건 / 크리덴셜 13%", "슬라이드는 **2,383건**, 현장 메모는 「20,383」으로 적혀 있으나 "
                                   "**슬라이드 기준이 맞다.** 크리덴셜 비중은 **녹취 13%**(메모 14%). "
                                   "원 출처가 슬라이드에 표기되지 않아 대외 인용 전 확인 필요"],
    ["87% API 공격", "슬라이드 기준. 녹취 ASR에는 97%로 들리나 **슬라이드가 정확**"],
    ["1,200+ / 1,000+ / 70+", "벤더 표기 수치. Entro 인수 발표의 수치와 일치한다 [1][2]"],
    ["「Agentic Fabric 도입 사례」", "**GA 1개월차**다. IGA 레퍼런스(삼성바이오로직스 등)를 "
                                   "에이전트 거버넌스 레퍼런스로 대체해 듣지 말 것"],
], widths=[4.0, 13.0], first_col_bold=True)

hr(doc)
doc.add_heading("출처", level=2)
sources = [
    "[1] SailPoint Acquires Entro Security to Strengthen Automated Machine Identity and Credential Lifecycle Management — NHI Mgmt Group  https://nhimg.org/nhi-news/sailpoint-acquires-entro-security-non-human-identity",
    "[2] SailPoint acquires Entro to bolster non-human identity — FinTech Global  https://fintech.global/2026/06/19/sailpoint-acquires-entro-to-bolster-non-human-identity/",
    "[3] SailPoint's Entro acquisition signals a shift in NHI governance — NHI Mgmt Group  https://nhimg.org/articles/sailpoints-entro-acquisition-signals-a-shift-in-nhi-governance/",
    "[4] SailPoint Virtual Appliances — SailPoint 공식 문서  https://documentation.sailpoint.com/saas/help/va/index.html",
    "[5] SaaS Connectors / Connector Architecture — SailPoint 공식 문서  https://documentation.sailpoint.com/connectors/isc/landingpages/help/landingpages/saas_connectors.html",
    "[6] SailPoint Agentic Fabric expands identity governance to autonomous AI agents — Help Net Security  https://www.helpnetsecurity.com/2026/05/11/sailpoint-agentic-fabric-expands-identity-governance-to-autonomous-ai-agents/",
    "[7] SailPoint Launches Agentic Fabric to Secure AI Identities Across the Enterprise (2026-05-11)  https://investor.sailpoint.com/news-releases/news-release-details/sailpoint-launches-agentic-fabric-secure-ai-identities-across",
    "[8] SailPoint 고객 사례 페이지  https://www.sailpoint.com/customers",
    "[9] Toyota Motor Europe 고객 사례 — SailPoint  https://www.sailpoint.com/ko/customers/toyota-motor-europe",
    "[10] Entro Security — Non-Human Identity and Secrets Security (Microsoft Marketplace)  https://marketplace.microsoft.com/en-us/product/entro-security.entro-security-marketplace?tab=overview",
    "[11] SailPoint eliminates blind spots with unified protection across human, non-human, and AI agent identities (2026-08-04, GA)  https://www.globenewswire.com/news-release/2026/08/04/3338366/0/en/SailPoint-eliminates-blind-spots-with-unified-protection-across-human-non-human-and-AI-agent-identities.html",
    "[12] Secure Agentic AI — SailPoint 제품 페이지  https://www.sailpoint.com/products/agentic-ai-security",
    "[13] SailPoint Customer Agreements Definitions and Additional Terms  https://documentation.sailpoint.com/main_landing_page/customer_agreements.html",
    "[14] OpenAI Says Its AI Models Escaped Sandbox, Targeted Hugging Face to Cheat Benchmark — The Hacker News  https://thehackernews.com/2026/07/openai-says-its-own-ai-models-escaped.html",
    "[15] OpenAI's accidental cyberattack against Hugging Face — Simon Willison (2026-07-22)  https://simonwillison.net/2026/Jul/22/openai-cyberattack/",
    "[16] How an AI Escaped Its Sandbox and Hacked Hugging Face — Better Stack  https://betterstack.com/community/guides/ai/openai-hugging-face/",
    "[17] SailPoint Research Highlights Rapid AI Agent Adoption (2025-05-28)  https://investor.sailpoint.com/news-releases/news-release-details/sailpoint-research-highlights-rapid-ai-agent-adoption-driving",
    "[18] AI agents: The new attack surface — SailPoint / Dimensional Research  https://www.sailpoint.com/press-releases/sailpoint-ai-agent-adoption-report",
]
for s in sources:
    para(doc, s, size=8.5, color=MUTED, space_after=2, indent=0.3)

out = "/home/user/Aistudy/share/03_SailPoint_AI에이전트-NHI거버넌스_세션정리.docx"
doc.save(out)
print("saved:", out)
